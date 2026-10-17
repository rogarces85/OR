"""Exportación de reportes: XLSX consolidado y PDF del informe calendario.

Cubre T7.2 (F-20, P-3) y T7.3 (F-21, P-3). El enlace ``mailto`` del aviso de
morosidad vive en :mod:`app.services.correo` (T7.4). El caso "exportación sin
datos" (sección 9 de SPECS) se resuelve en el router: si el periodo no tiene
movimientos se informa al usuario en pantalla y no se descarga un archivo vacío.
"""

from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.core.formato import formatear_clp
from app.core.periodos import MESES_NOMBRES
from app.models.membresia import ESTADO_PARCIAL
from app.models.miembro import CATEGORIA_ACTIVO, CATEGORIA_MEMBRESIA
from app.services.reportes import (
    InformeCalendario,
    MesInforme,
    ReporteConsolidado,
)

# Tipos MIME usados por los endpoints de descarga.
CONTENIDO_XLSX = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
CONTENIDO_PDF = "application/pdf"

_FORMATO_CLP = "#,##0"
_COLOR_MARCA = "1F2937"
_COLOR_GRIS = "F3F4F6"

# Categoría interna del miembro -> texto mostrado en los documentos.
_CATEGORIAS_TEXTO: dict[str, str] = {
    CATEGORIA_ACTIVO: "Alumno (Activo)",
    CATEGORIA_MEMBRESIA: "Membresía",
}


# --------------------------------------------------------------------------- #
# Utilidades comunes
# --------------------------------------------------------------------------- #


def _fecha_texto(valor: date) -> str:
    """Formatea una fecha como ``dd/mm/aaaa``."""
    return valor.strftime("%d/%m/%Y")


def _etiqueta_categoria(categoria: str) -> str:
    """Traduce la categoría interna del miembro a un texto legible."""
    return _CATEGORIAS_TEXTO.get(categoria, categoria)


# --------------------------------------------------------------------------- #
# XLSX del reporte consolidado (T7.2)
# --------------------------------------------------------------------------- #


def exportar_reporte_xlsx(reporte: ReporteConsolidado) -> bytes:
    """Genera el XLSX del reporte consolidado con los mismos totales (CA-11)."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Reporte consolidado"

    hoja["A1"] = "Osorno Runners - Reporte consolidado"
    hoja["A1"].font = Font(size=14, bold=True)
    hoja["A2"] = (
        f"Periodo: {_fecha_texto(reporte.desde)} a {_fecha_texto(reporte.hasta)}"
    )
    hoja["A2"].font = Font(italic=True, color="6B7280")

    fila = 4
    hoja.cell(row=fila, column=1, value="Resumen").font = Font(bold=True)
    fila += 1
    for etiqueta, valor in (
        ("Total ingresos", reporte.total_ingresos_clp),
        ("Total egresos", reporte.total_egresos_clp),
        ("Balance", reporte.neto_clp),
    ):
        hoja.cell(row=fila, column=1, value=etiqueta)
        celda = hoja.cell(row=fila, column=2, value=valor)
        celda.number_format = _FORMATO_CLP
        fila += 1

    fila += 1
    fila = _escribir_bloque(
        hoja,
        fila,
        f"Ingresos ({len(reporte.ingresos)})",
        ("Fecha", "Miembro", "Medio", "Monto (CLP)"),
        [
            (
                _fecha_texto(pago.fecha_pago),
                pago.miembro.nombre_completo if pago.miembro else "",
                pago.medio,
                int(pago.monto_clp),
            )
            for pago in reporte.ingresos
        ],
        columnas_monetarias=(4,),
        texto_sin_datos="Sin ingresos en el periodo.",
    )

    fila += 1
    _escribir_bloque(
        hoja,
        fila,
        f"Egresos ({len(reporte.egresos)})",
        ("Fecha", "Concepto", "Categoría", "Monto (CLP)"),
        [
            (
                _fecha_texto(egreso.fecha),
                egreso.concepto,
                egreso.categoria.nombre if egreso.categoria else "",
                int(egreso.monto_clp),
            )
            for egreso in reporte.egresos
        ],
        columnas_monetarias=(4,),
        texto_sin_datos="Sin egresos en el periodo.",
    )

    for columna, ancho in ((1, 16), (2, 44), (3, 20), (4, 16)):
        hoja.column_dimensions[
            hoja.cell(row=1, column=columna).column_letter
        ].width = ancho

    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def _escribir_bloque(
    hoja: Worksheet,
    fila: int,
    titulo: str,
    encabezados: tuple[str, ...],
    filas: list[tuple[object, ...]],
    *,
    columnas_monetarias: tuple[int, ...] = (),
    texto_sin_datos: str = "",
) -> int:
    """Escribe un título, su encabezado y sus filas; devuelve la fila libre."""
    hoja.cell(row=fila, column=1, value=titulo).font = Font(bold=True)
    fila += 1
    for columna, encabezado in enumerate(encabezados, start=1):
        celda = hoja.cell(row=fila, column=columna, value=encabezado)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=_COLOR_MARCA)
        celda.alignment = Alignment(
            horizontal="right" if columna in columnas_monetarias else "left"
        )
    fila += 1

    if not filas:
        hoja.cell(row=fila, column=1, value=texto_sin_datos or "Sin datos.")
        return fila + 1

    for datos in filas:
        for columna, valor in enumerate(datos, start=1):
            celda = hoja.cell(row=fila, column=columna, value=valor)
            if columna in columnas_monetarias:
                celda.number_format = _FORMATO_CLP
                celda.alignment = Alignment(horizontal="right")
        fila += 1
    return fila


# --------------------------------------------------------------------------- #
# PDF del informe calendario (T7.3)
# --------------------------------------------------------------------------- #


def exportar_informe_pdf(informe: InformeCalendario) -> bytes:
    """Genera el PDF del informe calendario con estados y montos (CA-12)."""
    buffer = BytesIO()
    documento = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title=(
            f"Informe calendario {informe.anio} - "
            f"{informe.miembro.nombre_completo}"
        ),
        author="Osorno Runners",
    )
    estilos = getSampleStyleSheet()
    elementos = [
        Paragraph("Osorno Runners", estilos["Title"]),
        Paragraph(f"Informe calendario {informe.anio}", estilos["Heading2"]),
        Paragraph(
            f"{informe.miembro.nombre_completo} · "
            f"{_etiqueta_categoria(informe.miembro.categoria)} · "
            f"{informe.miembro.email}",
            estilos["Normal"],
        ),
        Spacer(1, 6 * mm),
        _tabla_totales(informe),
        Spacer(1, 6 * mm),
        _tabla_calendario(informe),
        Spacer(1, 4 * mm),
        Paragraph(
            f"Generado el {_fecha_texto(date.today())}. "
            "Documento informativo del estado de cuotas del periodo.",
            estilos["Italic"],
        ),
    ]
    documento.build(elementos)
    return buffer.getvalue()


def _tabla_totales(informe: InformeCalendario) -> Table:
    """Tabla resumen con los totales del año (debe coincidir con la pantalla)."""
    datos = [
        ["Total del año", "Pagado", "Pendiente"],
        [
            f"$ {formatear_clp(informe.total_monto_clp)}",
            f"$ {formatear_clp(informe.total_pagado_clp)}",
            f"$ {formatear_clp(informe.total_pendiente_clp)}",
        ],
    ]
    tabla = Table(datos, colWidths=[60 * mm, 60 * mm, 60 * mm])
    tabla.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{_COLOR_GRIS}")),
                ("FONTNAME", (0, 0), (-1, 1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 10),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    return tabla


def _tabla_calendario(informe: InformeCalendario) -> Table:
    """Tabla con el estado y los montos de cada mes del año."""
    filas: list[list[str]] = [
        ["Mes", "Estado", "Monto (CLP)", "Pagado (CLP)", "Saldo (CLP)"]
    ]
    for mes in informe.meses:
        filas.append(
            [
                MESES_NOMBRES[mes.mes - 1],
                _etiqueta_estado(mes),
                _monto(mes.monto_clp, mes.tiene_cuota),
                _monto(mes.pagado_clp, mes.tiene_cuota),
                _monto(_saldo_mes(mes), mes.tiene_cuota),
            ]
        )
    filas.append(
        [
            "Total",
            "",
            f"$ {formatear_clp(informe.total_monto_clp)}",
            f"$ {formatear_clp(informe.total_pagado_clp)}",
            f"$ {formatear_clp(informe.total_pendiente_clp)}",
        ]
    )

    tabla = Table(filas, colWidths=[32 * mm, 30 * mm, 40 * mm, 40 * mm, 40 * mm])
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{_COLOR_MARCA}")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor(f"#{_COLOR_GRIS}")),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for indice, mes in enumerate(informe.meses, start=1):
        estilo.append(
            ("TEXTCOLOR", (1, indice), (1, indice), _color_estado(mes))
        )
    tabla.setStyle(TableStyle(estilo))
    return tabla


def _etiqueta_estado(mes: MesInforme) -> str:
    """Etiqueta del estado de un mes, igual que en la pantalla (UI-06)."""
    if not mes.tiene_cuota:
        return "Sin cuota"
    if mes.es_exenta:
        return "Exenta"
    if mes.es_pagada:
        return "Pagada"
    if mes.estado == ESTADO_PARCIAL:
        return "Parcial"
    return "Pendiente"


def _color_estado(mes: MesInforme) -> colors.Color:
    """Color del texto del estado según la semántica de morosidad (RN-21)."""
    if not mes.tiene_cuota or mes.es_exenta:
        return colors.HexColor("#6B7280")
    if mes.es_pagada:
        return colors.HexColor("#15803D")
    if mes.estado == ESTADO_PARCIAL:
        return colors.HexColor("#B45309")
    return colors.HexColor("#B91C1C")


def _saldo_mes(mes: MesInforme) -> int:
    """Saldo pendiente del mes (0 si está pagada, exenta o sin cuota)."""
    if mes.cuota is None or mes.es_pagada or mes.es_exenta:
        return 0
    return max(mes.monto_clp - mes.pagado_clp, 0)


def _monto(valor: int, tiene_cuota: bool) -> str:
    """Formatea un monto o un guion cuando el mes no tiene cuota."""
    if not tiene_cuota:
        return "-"
    return f"$ {formatear_clp(valor)}"
