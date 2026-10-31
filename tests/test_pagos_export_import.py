"""Tests para exportación e importación masiva de pagos (T5.17, T5.18)."""

from __future__ import annotations

import io
from datetime import date

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.dependencies import get_db
from app.main import app


@pytest.fixture()
def cliente(db):
    """TestClient con la sesión de prueba inyectada como dependencia."""

    def _get_db():
        yield db

    app.dependency_overrides[get_db] = _get_db
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


def _sesion_admin(cliente, usuario_factory):
    usuario_factory(email="admin@example.com", password="Clave1234")
    respuesta = cliente.get("/ingresar")
    assert respuesta.status_code == 200
    token = respuesta.text.split('name="csrf_token" value="')[1].split('"')[0]
    respuesta = cliente.post(
        "/ingresar",
        data={
            "email": "admin@example.com",
            "password": "Clave1234",
            "csrf_token": token,
        },
        follow_redirects=False,
    )
    assert respuesta.status_code == 303


def _crear_miembro_servicio(db, servicio_membresias, email, nombre):
    """Crea un miembro usando el servicio para pruebas."""
    return servicio_membresias.crear_miembro(
        db,
        nombre_completo=nombre,
        email=email,
        categoria="activo",
        fecha_incorporacion=date(2025, 1, 1),
    )


def test_exportar_pagos_devuelve_xlsx(cliente, usuario_factory, db, miembro_factory):
    """La exportación devuelve un archivo XLSX válido con los pagos."""
    from app.services import membresias as servicio_membresias

    _sesion_admin(cliente, usuario_factory)
    miembro = _crear_miembro_servicio(db, servicio_membresias, "exportar@example.com", "Para Exportar")

    # Registrar un pago para que exista algo que exportar
    from app.services import pagos as servicio_pagos
    servicio_pagos.registrar_pago(
        db,
        miembro=miembro,
        monto_clp=15000,
        anio=2026,
        meses=[10],
        fecha_pago=date(2026, 10, 15),
        medio="transferencia",
        observaciones="Pago de prueba",
    )

    respuesta = cliente.get("/pagos/exportar")
    assert respuesta.status_code == 200
    assert respuesta.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment; filename=" in respuesta.headers["content-disposition"]

    # Verificar que el archivo XLSX es válido y contiene los datos
    wb = load_workbook(io.BytesIO(respuesta.content))
    ws = wb.active
    filas = list(ws.iter_rows(values_only=True))
    assert len(filas) >= 2  # encabezado + al menos un pago
    assert filas[0] == (
        "Fecha", "Miembro", "Email", "RUT", "Medio de pago", "Monto (CLP)", "Observaciones"
    )
    # Buscar la fila de nuestro pago
    encontrado = False
    for fila in filas[1:]:
        if fila[1] == "Para Exportar" and fila[2] == "exportar@example.com":
            assert fila[4] == "transferencia"
            assert fila[5] == 15000
            assert fila[6] == "Pago de prueba"
            encontrado = True
    assert encontrado, "El pago exportado no se encontró en el XLSX"


def test_importar_pagos_desde_xlsx(cliente, usuario_factory, db, miembro_factory):
    """La importación desde XLSX crea los pagos correctamente."""
    from app.services import membresias as servicio_membresias

    _sesion_admin(cliente, usuario_factory)
    miembro = _crear_miembro_servicio(db, servicio_membresias, "importar@example.com", "Para Importar")

    # Crear archivo XLSX en memoria
    wb = load_workbook(io.BytesIO()) if False else None
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["identificador", "monto", "fecha_pago", "medio", "observaciones"])
    ws.append(["importar@example.com", 20000, "2026-10-20", "transferencia", "Pago importado"])
    ws.append([miembro.email, 10000, "2026-10-21", "efectivo", "Pago en efectivo"])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    respuesta = cliente.post(
        "/pagos/importar",
        files={"archivo": ("pagos_test.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert respuesta.status_code == 200
    assert "2 pago(s) importado(s)" in respuesta.text

    # Verificar que los pagos se crearon en la base de datos
    from app.repositories import pago as repo_pago
    pagos = repo_pago.listar(db)
    importar_pagos = [p for p in pagos if p.miembro.email == "importar@example.com"]
    assert len(importar_pagos) == 2


def test_importar_pagos_desde_csv(cliente, usuario_factory, db, miembro_factory):
    """La importación desde CSV crea los pagos correctamente."""
    from app.services import membresias as servicio_membresias

    _sesion_admin(cliente, usuario_factory)
    miembro = _crear_miembro_servicio(db, servicio_membresias, "csv@example.com", "Para CSV")

    contenido_csv = (
        "identificador,monto,fecha_pago,medio,observaciones\n"
        "csv@example.com,25000,2026-10-22,transferencia,Pago CSV\n"
    )

    respuesta = cliente.post(
        "/pagos/importar",
        files={"archivo": ("pagos_test.csv", contenido_csv.encode("utf-8"), "text/csv")},
    )
    assert respuesta.status_code == 200
    assert "1 pago(s) importado(s)" in respuesta.text

    from app.repositories import pago as repo_pago
    pagos = repo_pago.listar(db)
    csv_pagos = [p for p in pagos if p.miembro.email == "csv@example.com"]
    assert len(csv_pagos) == 1
    assert csv_pagos[0].monto_clp == 25000


def test_importar_pagos_miembro_no_existente(cliente, usuario_factory, db):
    """La importación falla si el miembro no existe."""
    _sesion_admin(cliente, usuario_factory)

    contenido_xlsx = (
        "identificador,monto,fecha_pago,medio,observaciones\n"
        "noexiste@example.com,10000,2026-10-22,transferencia,Pago\n"
    )

    # Crear XLSX simple
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["identificador", "monto", "fecha_pago", "medio", "observaciones"])
    ws.append(["noexiste@example.com", 10000, "2026-10-22", "transferencia", "Pago"])
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    respuesta = cliente.post(
        "/pagos/importar",
        files={"archivo": ("pagos_test.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert respuesta.status_code == 200
    assert "1 pago(s) no se pudieron importar" in respuesta.text
