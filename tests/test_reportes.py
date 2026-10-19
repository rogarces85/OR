"""Informe calendario y exportaciones XLSX/PDF (CA-11, CA-14)."""

from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from app.models.membresia import (
    ESTADO_EXENTA,
    ESTADO_PAGADA,
    ESTADO_PENDIENTE,
    NIVEL_ACTIVA,
)
from app.repositories import finanzas as repo_finanzas
from app.repositories import pago as repo_pago
from app.services import exportacion as servicio_exportacion
from app.services import reportes as servicio_reportes


def test_informe_calendario_estados_y_totales(
    db, niveles, miembro_factory, cuota_factory
):
    """CA-14: el informe distingue pagadas, pendientes y exentas."""
    miembro = miembro_factory()
    activa = niveles[NIVEL_ACTIVA]
    cuota_factory(
        miembro, activa, 2026, 1, estado=ESTADO_PAGADA, pagado=15_000
    )
    cuota_factory(miembro, activa, 2026, 2, estado=ESTADO_PENDIENTE)
    cuota_factory(miembro, activa, 2026, 3, estado=ESTADO_EXENTA)

    informe = servicio_reportes.informe_calendario(db, miembro, 2026)

    assert len(informe.meses) == 12
    assert informe.meses[0].es_pagada is True
    assert informe.meses[1].estado == ESTADO_PENDIENTE
    assert informe.meses[2].es_exenta is True
    assert informe.meses[3].tiene_cuota is False

    assert informe.total_monto_clp == 45_000
    assert informe.total_pagado_clp == 15_000
    assert informe.total_pendiente_clp == 15_000


def test_exportacion_xlsx_contiene_los_mismos_totales(
    db, categorias, miembro_factory
):
    """CA-11: el XLSX descargado contiene los totales mostrados en pantalla."""
    miembro = miembro_factory()
    for monto, dia in ((30_000, 5), (10_000, 6)):
        repo_pago.crear(
            db,
            miembro_id=miembro.id,
            monto_clp=monto,
            fecha_pago=date(2026, 7, dia),
        )
    repo_finanzas.crear_egreso(
        db,
        monto_clp=12_000,
        concepto="Egreso julio",
        categoria_egreso_id=categorias["Otros"].id,
        fecha=date(2026, 7, 8),
    )
    db.commit()

    reporte = servicio_reportes.reporte_consolidado(
        db, desde=date(2026, 7, 1), hasta=date(2026, 7, 31)
    )
    contenido = servicio_exportacion.exportar_reporte_xlsx(reporte)

    assert contenido.startswith(b"PK")
    libro = load_workbook(BytesIO(contenido))
    hoja = libro.active
    resumen = {
        hoja.cell(row=fila, column=1).value: hoja.cell(row=fila, column=2).value
        for fila in range(1, hoja.max_row + 1)
        if hoja.cell(row=fila, column=1).value
    }
    assert resumen["Total ingresos"] == reporte.total_ingresos_clp == 40_000
    assert resumen["Total egresos"] == reporte.total_egresos_clp == 12_000
    assert resumen["Balance"] == reporte.neto_clp == 28_000


def test_exportacion_xlsx_sin_movimientos(db):
    reporte = servicio_reportes.reporte_consolidado(
        db, desde=date(2026, 9, 1), hasta=date(2026, 9, 30)
    )
    contenido = servicio_exportacion.exportar_reporte_xlsx(reporte)

    libro = load_workbook(BytesIO(contenido))
    valores = [
        celda.value
        for fila in libro.active.iter_rows()
        for celda in fila
        if celda.value
    ]
    assert "Sin ingresos en el periodo." in valores
    assert "Sin egresos en el periodo." in valores


def test_exportacion_pdf_del_informe(
    db, niveles, miembro_factory, cuota_factory
):
    """CA-14: el PDF refleja el estado y los montos del informe."""
    miembro = miembro_factory()
    cuota_factory(
        miembro,
        niveles[NIVEL_ACTIVA],
        2026,
        1,
        estado=ESTADO_PAGADA,
        pagado=15_000,
    )
    cuota_factory(miembro, niveles[NIVEL_ACTIVA], 2026, 2)
    informe = servicio_reportes.informe_calendario(db, miembro, 2026)

    contenido = servicio_exportacion.exportar_informe_pdf(informe)

    assert contenido.startswith(b"%PDF")
    assert len(contenido) > 1_000
